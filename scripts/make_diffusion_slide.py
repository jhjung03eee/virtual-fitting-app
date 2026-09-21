"""'diffusion 모델이란? / 왜 채택했나' 단독 슬라이드를 만든다.

발표 본 파일은 사용자가 PowerPoint에서 직접 고치고 있어서 건드리지 않고, 같은 디자인의
한 장짜리 파일로 만들어 복사해 넣게 한다.

왼쪽의 4단계 그림은 실제 합성 과정의 캡처가 아니라, 목업 사진에 노이즈를 섞어
"노이즈에서 점점 선명해진다"는 개념만 보여 주는 설명용 이미지다.

오른쪽 GAN/diffusion 결과 비교 사진은 IDM-VTON 논문(Choi et al., ECCV 2024) Fig. 4의
VITON-HD 두 번째 줄에서 잘라 쓴다(HR-VITON = GAN, IDM-VTON = diffusion). 원본 그림은
저작물이라 gitignore 폴더 data/samples/idm_vton_fig4.png에 두고, 슬라이드에 출처를 적는다.

    python scripts/make_diffusion_slide.py
"""
import os
import sys

from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_proposal_deck import (  # noqa: E402  디자인 토큰·도구를 발표 자료와 공유
    BG, COL_W, FG, FG_DIM, FG_SUB, GREEN, ROOT, SLIDE_H, SLIDE_W, M,
    F_LIGHT, F_MED, F_REG, F_THIN,
    _emu, _eyebrow, _lines, _picture, _rule, _textbox, _timed_notes, _vrule,
)

SOURCE_PHOTO = os.path.join(ROOT, 'design', 'fitting_result.jpg')
STEP_DIR = os.path.join(ROOT, 'docs', 'mockups')
OUT_PATH = os.path.join(ROOT, 'docs', 'slide_why_diffusion.pptx')

STEPS = [('노이즈', 1.0), ('윤곽', 0.62), ('형태', 0.3), ('완성', 0.0)]
PAPER_FIG = os.path.join(ROOT, 'data', 'samples', 'idm_vton_fig4.png')
# (라벨, 그림 안 잘라낼 영역) — Fig. 4 (a) VITON-HD 두 번째 줄
FIG_CROPS = [
    ('입력 옷', (2, 539, 118, 694)),
    ('GAN · HR-VITON', (123, 379, 358, 695)),
    ('diffusion · IDM-VTON', (1323, 379, 1558, 695)),
]
FIG_SOURCE = '같은 옷을 입힌 결과  ·  출처: Choi et al., IDM-VTON (ECCV 2024) Fig. 4'
# (항목, GAN, diffusion, diffusion 쪽이 나은가)
COMPARE = [
    ('자연스러움', '어색함', '자연스러움', True),
    ('일반 사진', '약함', '강함', True),
    ('속도', '빠름', '느림', False),
]
NOTES = (
    '합성에는 diffusion 모델을 씁니다. 왼쪽 그림처럼 노이즈에서 시작해 조금씩 다듬어 사진을 '
    '완성하는 AI입니다.\n\n'
    '오른쪽은 같은 옷을 입힌 결과입니다. 예전 GAN 방식은 줄무늬가 사라지고 머리카락 주변이 '
    '뭉개졌지만, diffusion 방식은 옷의 무늬와 주름까지 자연스럽게 살아 있습니다.\n\n'
    '대신 속도는 느린데, 여러 장을 만들어 가장 자연스러운 한 장을 고르는 방식으로 보완합니다.'
)


def make_step_images():
    """목업 사진에 노이즈를 섞은 설명용 4단계 이미지를 만든다."""
    photo = Image.open(SOURCE_PHOTO).convert('RGB')
    w, h = photo.size
    noise = Image.merge('RGB', [Image.effect_noise((w, h), 110) for _ in range(3)])
    paths = []
    for i, (_, amount) in enumerate(STEPS):
        blurred = photo.filter(ImageFilter.GaussianBlur(radius=10 * amount))
        mixed = Image.blend(blurred, noise, amount)
        path = os.path.join(STEP_DIR, f'diffusion_step_{i + 1}.png')
        mixed.save(path)
        paths.append(path)
    return paths


def build(slide, step_paths):
    _eyebrow(slide, '합성 방식')
    tf = _textbox(slide, M, Inches(1.35), COL_W, Inches(0.7))
    _lines(tf, ['diffusion 모델을 선택한 이유'], 32, font=F_THIN)

    mid = _emu(SLIDE_W / 2)
    _vrule(slide, mid, Inches(2.5), Inches(4.3))

    # 왼쪽: diffusion 모델이란? — 한 줄 + 그림
    left_w = _emu(mid - M - Inches(0.45))
    tf = _textbox(slide, M, Inches(2.5), left_w, Inches(0.35))
    _lines(tf, ['diffusion 모델이란?'], 16, font=F_MED, color=FG)
    tf = _textbox(slide, M, Inches(2.95), left_w, Inches(0.35))
    _lines(tf, ['노이즈에서 조금씩 다듬어 사진을 만드는 AI'], 15, font=F_LIGHT, color=FG_SUB)

    gap = Inches(0.26)
    # 네 장 + 화살표 칸이 왼쪽 칸 폭을 넘지 않는 높이 (사진 비율 390:520)
    img_h = _emu((left_w - 3 * gap) / 4 * 520 / 390)
    img_top = Inches(3.75)
    x = M
    for i, ((label, _), path) in enumerate(zip(STEPS, step_paths)):
        pic = _picture(slide, path, x, img_top, img_h)
        tf = _textbox(slide, x, img_top + img_h + Inches(0.08), pic.width, Inches(0.3),
                      align=PP_ALIGN.CENTER)
        _lines(tf, [label], 11.5, font=F_REG, color=GREEN if i == len(STEPS) - 1 else FG_SUB,
               align=PP_ALIGN.CENTER)
        if i < len(STEPS) - 1:
            tf = _textbox(slide, x + pic.width, img_top + img_h / 2 - Inches(0.16), gap, Inches(0.32),
                          align=PP_ALIGN.CENTER)
            _lines(tf, ['→'], 13, font=F_LIGHT, color=FG_DIM, align=PP_ALIGN.CENTER)
        x = _emu(x + pic.width + gap)

    # 오른쪽: 기존 방식(GAN)과 비교 — 결과 사진 + 한 단어 표
    rx = _emu(mid + Inches(0.45))
    right_w = _emu(SLIDE_W - M - rx)
    tf = _textbox(slide, rx, Inches(2.5), right_w, Inches(0.35))
    _lines(tf, ['기존 방식(GAN)과 비교하면'], 16, font=F_MED, color=FG)
    tf = _textbox(slide, rx, Inches(2.92), right_w, Inches(0.3))
    _lines(tf, [FIG_SOURCE], 9.5, font=F_REG, color=FG_DIM)

    fig = Image.open(PAPER_FIG)
    fig_h = Inches(1.95)
    fig_gap = Inches(0.18)
    fig_top = Inches(3.3)
    x = rx
    for i, (label, box) in enumerate(FIG_CROPS):
        path = os.path.join(ROOT, 'data', 'samples', f'gan_vs_diffusion_{i + 1}.png')
        fig.crop(box).save(path)
        pic = _picture(slide, path, x, fig_top, fig_h)
        is_diffusion = i == len(FIG_CROPS) - 1
        tf = _textbox(slide, x, fig_top + fig_h + Inches(0.06), pic.width + Inches(0.2), Inches(0.28))
        _lines(tf, [label], 10.5, font=F_REG, color=GREEN if is_diffusion else FG_SUB)
        x = _emu(x + pic.width + fig_gap)

    label_w = Inches(1.3)
    col_w = _emu((right_w - label_w) / 2)
    row_h = Inches(0.3)
    y = Inches(5.72)
    for c, (name, color) in enumerate([('GAN', FG_SUB), ('diffusion', GREEN)]):
        tf = _textbox(slide, rx + label_w + col_w * c, y, col_w, row_h)
        _lines(tf, [name], 11.5, font=F_REG, color=color)
    _rule(slide, rx, y + row_h, right_w)
    y = _emu(y + row_h + Inches(0.06))
    for label, old, new, better in COMPARE:
        tf = _textbox(slide, rx, y, label_w, row_h)
        _lines(tf, [label], 11.5, font=F_MED, color=FG_DIM)
        tf = _textbox(slide, rx + label_w, y, col_w, row_h)
        _lines(tf, [old], 12, font=F_LIGHT, color=FG_SUB)
        tf = _textbox(slide, rx + label_w + col_w, y, col_w, row_h)
        _lines(tf, [new], 12, font=F_REG if better else F_LIGHT, color=FG if better else FG_SUB)
        y = _emu(y + row_h)

    _picture(slide, os.path.join(ROOT, 'docs', 'brand', 'fitcheck_wordmark_dim.png'), M,
             SLIDE_H - Inches(0.67), Inches(0.1))


def main():
    step_paths = make_step_images()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    build(slide, step_paths)
    # 단독 장이라 누적 시간은 의미가 없다. 이 장 소요 시간만 붙인다
    _, seconds = _timed_notes(NOTES, 0.0)
    slide.notes_slide.notes_text_frame.text = f'[약 {round(seconds / 5) * 5}초]\n\n{NOTES}'
    out = sys.argv[1] if len(sys.argv) > 1 else OUT_PATH  # 기본 파일이 PowerPoint에 열려 있을 때
    prs.save(out)
    print(f'슬라이드 생성 완료 (대본 약 {int(seconds)}초) -> {out}')


if __name__ == '__main__':
    main()
