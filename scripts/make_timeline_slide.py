"""'진행 계획' 타임라인 단독 슬라이드를 만든다.

발표 본 파일은 사용자가 PowerPoint에서 직접 고치고 있어서 건드리지 않고, 같은 디자인의
한 장짜리 파일로 만들어 복사해 넣게 한다. 월별 4단계 + 마감(계획서·중간·최종발표).

    python scripts/make_timeline_slide.py [저장.pptx]
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_proposal_deck import (  # noqa: E402  디자인 토큰·도구를 발표 자료와 공유
    BG, COL_W, FG, FG_DIM, FG_SUB, GREEN, LINE, ROOT, SLIDE_H, SLIDE_W, SURFACE, M,
    F_LIGHT, F_MED, F_REG, F_THIN,
    _emu, _eyebrow, _lines, _picture, _rect, _rule, _set_font, _textbox, _timed_notes,
)

OUT_PATH = os.path.join(ROOT, 'docs', 'slide_timeline.pptx')
CURRENT = 0  # 지금이 몇 번째 단계인지 (9월)

# (월, 단계, 할 일 3개, 마감)
PHASES = [
    ('9월', '조사 · 설계', ['관련 연구 · 모델 조사', '모델 후보 4종 선정', '앱 화면 설계'], '과제 계획서 제출'),
    ('10월', '모델 검증', ['후보 4종 같은 조건 비교', '휴대폰 사진으로 검증', '치수표 데이터 수집'], '주력 모델 확정'),
    ('11월', '구현 · 통합', ['사이즈 추천 로직 구현', '합성 + 추천 연결', '앱 화면 구현'], '중간발표 (11월 초)'),
    ('12월', '평가 · 발표', ['품질 개선', '사용자 평가', '시연 준비'], '최종발표 · 시연'),
]
NOTES = (
    '진행 계획입니다. 네 달을 네 단계로 나눴습니다.\n\n'
    '지금인 9월에는 관련 연구와 모델을 조사하고 계획서를 제출합니다. 10월에는 후보 모델 네 개를 '
    '같은 조건에서 비교해 주력 모델을 확정하고, 11월에는 사이즈 추천과 합성을 연결해 중간발표에서 '
    '보여 드립니다. 12월에는 품질을 개선하고 사용자 평가를 거쳐 최종발표에서 시연하겠습니다.'
)


def build(slide):
    _eyebrow(slide, '진행 계획')
    tf = _textbox(slide, M, Inches(1.35), COL_W, Inches(0.7))
    _lines(tf, ['언제까지 무엇을 끝내는가'], 32, font=F_THIN)

    col_w = _emu(COL_W / len(PHASES))
    month_y = Inches(2.45)
    axis_y = Inches(3.2)
    _rule(slide, M, axis_y, COL_W, color=LINE, weight=Pt(1.5))

    dot = Inches(0.18)
    for i, (month, phase, tasks, milestone) in enumerate(PHASES):
        x = _emu(M + col_w * i)
        inner_w = col_w - Inches(0.3)
        current = i == CURRENT

        tf = _textbox(slide, x, month_y, inner_w, Inches(0.6))
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = month
        _set_font(run, 26, font=F_THIN, color=FG)
        if current:
            tag = p.add_run()
            tag.text = '   지금'
            _set_font(tag, 11, font=F_MED, color=GREEN)

        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, _emu(x), _emu(axis_y - dot / 2), dot, dot)
        marker.fill.solid()
        marker.fill.fore_color.rgb = GREEN if current else FG_DIM
        marker.line.fill.background()
        marker.shadow.inherit = False

        tf = _textbox(slide, x, Inches(3.5), inner_w, Inches(0.4))
        _lines(tf, [phase], 17, font=F_REG, color=FG)

        tf = _textbox(slide, x, Inches(4.05), inner_w, Inches(1.2))
        _lines(tf, ['·  ' + t for t in tasks], 12.5, font=F_LIGHT, color=FG_SUB, spacing=1.15, gap=5)

        _rect(slide, x, Inches(5.5), inner_w, Inches(0.5), SURFACE, rounded=0.2)
        tf = _textbox(slide, x, Inches(5.5), inner_w, Inches(0.5), anchor=MSO_ANCHOR.MIDDLE,
                      align=PP_ALIGN.CENTER)
        _lines(tf, [milestone], 12.5, font=F_REG, color=GREEN if current else FG, align=PP_ALIGN.CENTER)

    _picture(slide, os.path.join(ROOT, 'docs', 'brand', 'fitcheck_wordmark_dim.png'), M,
             SLIDE_H - Inches(0.67), Inches(0.1))


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    build(slide)
    _, seconds = _timed_notes(NOTES, 0.0)
    slide.notes_slide.notes_text_frame.text = f'[약 {round(seconds / 5) * 5}초]\n\n{NOTES}'
    out = sys.argv[1] if len(sys.argv) > 1 else OUT_PATH
    prs.save(out)
    print(f'슬라이드 생성 완료 (대본 약 {int(seconds)}초) -> {out}')


if __name__ == '__main__':
    main()
